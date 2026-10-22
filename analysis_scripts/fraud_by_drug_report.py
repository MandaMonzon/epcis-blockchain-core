#!/usr/bin/env python3
"""
fraud_by_drug_report.py
========================
Extensao do relatorio do dia 14/09 (analysis/multi_scale_report.py), criada em
2026-09-28 a pedido da Amanda para responder aos seguintes requisitos:

  1. Restringir a analise as escalas com MAIS DE 50 medicamentos, ignorando
     d10_e100, d20_e100 e d50_e1000 (--min-drugs-scale, default 50).
  2. Calcular o PERCENTUAL DE FRAUDE por medicamento (drug_code + drug_name),
     nao a taxa de sucesso de transacao (SLA) que o relatorio original media.
     "Fraude" aqui = riskScore >= cutoff (default 90), numa escala 0-100 em
     que riskScore = soma dos scores das regras disparadas, clamped em 100
     (ver epcis-fraud-logistic-diversion/chaincode/RuleEngine.js:102). Esta
     definicao foi confirmada explicitamente por Amanda em 28/09/2026
     ("o calculo da fraude vem do score de risco") - NAO compara com o
     anomaly_type real do cenario, e' puramente sobre o score computado
     pela chaincode.
  3. Trazer os resultados quebrados POR RODADA (campo roundIndex, ja presente
     em cada registro de round_scores_raw/*.jsonl - nao precisa ser extraido
     do nome do arquivo).
  4. Deixar Drug Code e Drug Name visiveis ao lado do percentual de fraude.

IMPORTANTE (ponto de corte tem DOIS sentidos, nao confundir):
  - --min-drugs-scale (default 50): corte ESTRUTURAL - quais escalas entram
    no relatorio (numero de medicamentos da escala, ex.: d100_e1000 = 100).
  - --fraud-cutoff-pct (default 90): corte de SCORE - a partir de qual
    riskScore um evento e' contado como "fraude" no percentual por medicamento.

Reaproveita (via import) as funcoes de carregamento/juncao/filtragem de
multi_scale_report.py para garantir que a base de dados usada e' EXATAMENTE
a mesma do relatorio original de 14/09 - so a metrica de "abaixo do corte"
muda de txSuccess (SLA de infraestrutura) para riskScore (fraude).

Uso (rodar em node003, dentro de ~/Workspace/blockchain-research/analysis/):

  python3 fraud_by_drug_report.py \
      --runroot ~/Workspace/blockchain-research/RUN_2026-09-27_multiscale \
      --arcos-gz ~/arcos_raw.gz \
      --min-drugs-scale 50 \
      --fraud-cutoff-pct 90 \
      --output ~/Workspace/blockchain-research/RUN_2026-09-27_multiscale/fraud_report_gt50.xlsx

Saida: um .xlsx com:
  - "0. Resumo Fraude": uma linha por escala (>50 medicamentos), com
    percentual medio de fraude e quantos medicamentos tiveram fraude
    detectada.
  - "{escala} Evol": reaproveitada do relatorio original (score de risco
    medio ao longo do tempo/rodada, por medicamento) - ja existia e
    responde a "variacao do score ao longo do tempo".
  - "{escala} Fraude": uma linha por medicamento, com Drug Code, Drug Name,
    Total de Eventos, Score Medio, Eventos com Fraude e Fraude % - NOVO.
  - "{escala} Fraude x Rodada": mesma coisa, quebrada por rodada (r0, r1,
    r2, ...), com pivot e grafico de linha - NOVO, responde a "trazer
    resultados por cada rodada que executa".
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from results_collector import load_events, load_scenario, load_drug_names
from multi_scale_report import (
    load_round_scores,
    join_round_scores_to_drugs,
    filter_low_volume_drugs,
    build_score_evolution_table,
    DEFAULT_MIN_EVENTS_PER_DRUG,
)

DEFAULT_FRAUD_CUTOFF_PCT = 90.0
DEFAULT_MIN_DRUGS_SCALE = 50


def scale_drug_count(label, manifest):
    """Numero de medicamentos da escala. Prioriza manifest['max_drugs'];
    fallback: extrai do nome da pasta (ex.: 'd100_e1000' -> 100)."""
    if manifest and manifest.get("max_drugs") not in (None, "", "?"):
        try:
            return int(manifest["max_drugs"])
        except (TypeError, ValueError):
            pass
    m = re.match(r"d(\d+)_e\d+", label)
    return int(m.group(1)) if m else None


def build_fraud_cutoff_table(filtered, cutoff_pct):
    """Analogo a build_sla_cutoff_table do relatorio original, mas baseado em
    riskScore (fraude) em vez de txSuccess (SLA de infraestrutura)."""
    groups = defaultdict(lambda: {"total": 0, "flagged": 0, "sum_score": 0.0, "name": ""})
    for r in filtered:
        score = r.get("riskScore")
        if score is None:
            continue
        score = float(score)
        g = groups[r["drug_code"]]
        g["total"] += 1
        g["sum_score"] += score
        if score >= cutoff_pct:
            g["flagged"] += 1
        g["name"] = r.get("drug_name", r["drug_code"])

    rows = []
    for drug_code, g in sorted(groups.items(), key=lambda kv: kv[1]["total"], reverse=True):
        if g["total"] == 0:
            continue
        fraud_pct = 100.0 * g["flagged"] / g["total"]
        rows.append({
            "drug_code": drug_code,
            "drug_name": g["name"],
            "total_events": g["total"],
            "avg_risk_score": round(g["sum_score"] / g["total"], 2),
            "n_flagged": g["flagged"],
            "fraud_pct": round(fraud_pct, 1),
        })
    return rows


def build_fraud_cutoff_table_by_round(filtered, cutoff_pct):
    """Mesma logica de build_fraud_cutoff_table, quebrada por rodada
    (roundIndex - campo ja presente em cada registro do round_scores_raw)."""
    groups = defaultdict(lambda: {"total": 0, "flagged": 0, "sum_score": 0.0, "name": ""})
    for r in filtered:
        score = r.get("riskScore")
        if score is None:
            continue
        score = float(score)
        key = (r.get("roundIndex"), r["drug_code"])
        g = groups[key]
        g["total"] += 1
        g["sum_score"] += score
        if score >= cutoff_pct:
            g["flagged"] += 1
        g["name"] = r.get("drug_name", r["drug_code"])

    rows = []
    for (round_idx, drug_code), g in sorted(
        groups.items(), key=lambda kv: (kv[0][0] is None, kv[0][0], -kv[1]["total"])
    ):
        if g["total"] == 0:
            continue
        fraud_pct = 100.0 * g["flagged"] / g["total"]
        rows.append({
            "round": round_idx,
            "drug_code": drug_code,
            "drug_name": g["name"],
            "total_events": g["total"],
            "avg_risk_score": round(g["sum_score"] / g["total"], 2),
            "n_flagged": g["flagged"],
            "fraud_pct": round(fraud_pct, 1),
        })
    return rows


def compute_fraud_scale_summary(fraud_rows, label, n_drugs_scale):
    if not fraud_rows:
        return {"label": label, "medicamentos_escala": n_drugs_scale, "n_drugs_avaliados": 0,
                "fraud_pct_medio": None, "n_drugs_com_fraude": 0}
    n_above = sum(1 for r in fraud_rows if r["n_flagged"] > 0)
    avg = sum(r["fraud_pct"] for r in fraud_rows) / len(fraud_rows)
    return {
        "label": label,
        "medicamentos_escala": n_drugs_scale,
        "n_drugs_avaliados": len(fraud_rows),
        "fraud_pct_medio": round(avg, 1),
        "n_drugs_com_fraude": n_above,
    }


def write_workbook(summary_rows, per_scale_data, cutoff_pct, min_drugs_scale, output_path):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "0. Resumo Fraude"
    ws0.append([
        f"Escopo: escalas com mais de {int(min_drugs_scale)} medicamentos "
        f"(ignora d10/d20/d50). Corte de fraude: riskScore >= {int(cutoff_pct)} "
        f"(escala 0-100)."
    ])
    ws0.append([])
    headers0 = ["Escala", "Medicamentos (escala)", "Medicamentos Avaliados",
                "Fraude % Media", "Medicamentos com Fraude Detectada"]
    ws0.append(headers0)
    for cell in ws0[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C0504D")
    for row in summary_rows:
        ws0.append([row["label"], row["medicamentos_escala"], row["n_drugs_avaliados"],
                    row["fraud_pct_medio"], row["n_drugs_com_fraude"]])
    for col in ws0.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws0.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    for label, data in per_scale_data.items():
        evolution_rows = data["evolution"]
        fraud_rows = data["fraud"]
        fraud_by_round = data["fraud_by_round"]

        # --- Aba de evolucao de score (reaproveitada do relatorio original) ---
        ws1 = wb.create_sheet(f"{label} Evol"[:31])
        ws1.append(["Round", "Drug Code", "Drug Name", "Event Type", "Avg Risk Score", "N Events"])
        for cell in ws1[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
        if not evolution_rows:
            ws1.append(["(sem dados)"])
        else:
            for r in evolution_rows:
                ws1.append([r["round"], r["drug_code"], r["drug_name"], r["event_type"],
                            r["avg_risk_score"], r["n_events"]])
        for col in ws1.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws1.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

        # --- Aba de fraude por medicamento (NOVA) ---
        ws2 = wb.create_sheet(f"{label} Fraude"[:31])
        ws2.append([f"Cutoff de fraude: riskScore >= {int(cutoff_pct)}."])
        ws2.append([])
        ws2.append(["Drug Code", "Drug Name", "Total Events", "Avg Risk Score",
                    f"Eventos Fraude (>={int(cutoff_pct)})", "Fraude %"])
        for cell in ws2[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C0504D")
        if not fraud_rows:
            ws2.append(["(sem dados)"])
        else:
            for r in fraud_rows:
                ws2.append([r["drug_code"], r["drug_name"], r["total_events"],
                            r["avg_risk_score"], r["n_flagged"], r["fraud_pct"]])
            n_com_fraude = sum(1 for r in fraud_rows if r["n_flagged"] > 0)
            ws2.append([])
            ws2.append([f"Medicamentos com fraude detectada (score>={int(cutoff_pct)} em pelo menos 1 evento):",
                        n_com_fraude, f"de {len(fraud_rows)} avaliados"])
        for col in ws2.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws2.column_dimensions[col[0].column_letter].width = min(max_len + 2, 45)

        # --- Aba de fraude por medicamento x rodada (NOVA) ---
        ws3 = wb.create_sheet(f"{label} Fraude x Rodada"[:31])
        ws3.append(["Round", "Drug Code", "Drug Name", "Total Events", "Avg Risk Score",
                    "Eventos Fraude", "Fraude %"])
        for cell in ws3[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
        if not fraud_by_round:
            ws3.append(["(sem dados)"])
        else:
            for r in fraud_by_round:
                ws3.append([r["round"], r["drug_code"], r["drug_name"], r["total_events"],
                            r["avg_risk_score"], r["n_flagged"], r["fraud_pct"]])

            rounds = sorted({r["round"] for r in fraud_by_round if r["round"] is not None})
            drugs = sorted({(r["drug_code"], r["drug_name"]) for r in fraud_by_round})
            if len(rounds) >= 2 and drugs:
                pivot_start = ws3.max_row + 3
                ws3.cell(row=pivot_start, column=1,
                         value="Pivot (fraude % por rodada, por medicamento)").font = Font(bold=True)
                header_row = pivot_start + 1
                ws3.cell(row=header_row, column=1, value="Round")
                for j, (code, name) in enumerate(drugs, start=2):
                    ws3.cell(row=header_row, column=j, value=name or code)

                lookup = {(r["round"], r["drug_code"]): r["fraud_pct"] for r in fraud_by_round}
                for i, rnd in enumerate(rounds, start=1):
                    r_row = header_row + i
                    ws3.cell(row=r_row, column=1, value=rnd)
                    for j, (code, name) in enumerate(drugs, start=2):
                        val = lookup.get((rnd, code))
                        if val is not None:
                            ws3.cell(row=r_row, column=j, value=val)

                chart = LineChart()
                chart.title = f"Fraude % por medicamento e rodada - {label}"
                chart.x_axis.title = "Rodada"
                chart.y_axis.title = "Fraude %"
                data_ref = Reference(ws3, min_col=2, max_col=1 + len(drugs),
                                      min_row=header_row, max_row=header_row + len(rounds))
                cats = Reference(ws3, min_col=1, max_col=1,
                                  min_row=header_row + 1, max_row=header_row + len(rounds))
                chart.add_data(data_ref, titles_from_data=True)
                chart.set_categories(cats)
                ws3.add_chart(chart, f"A{header_row + len(rounds) + 3}")
        for col in ws3.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws3.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--runroot", required=True)
    parser.add_argument("--arcos-gz", default=None)
    parser.add_argument("--min-events-per-drug", type=int, default=DEFAULT_MIN_EVENTS_PER_DRUG)
    parser.add_argument("--fraud-cutoff-pct", type=float, default=DEFAULT_FRAUD_CUTOFF_PCT)
    parser.add_argument("--min-drugs-scale", type=int, default=DEFAULT_MIN_DRUGS_SCALE,
                         help="Ignora escalas com numero de medicamentos <= este valor (default 50)")
    parser.add_argument("--output", default=None)
    args = parser.parse_args()

    runroot = Path(args.runroot)
    output_path = Path(args.output) if args.output else runroot / "fraud_report_gt50.xlsx"
    drug_names = load_drug_names(args.arcos_gz) if args.arcos_gz else {}

    scale_dirs = sorted(d for d in runroot.iterdir() if d.is_dir())
    summary_rows = []
    per_scale_data = {}

    for scale_dir in scale_dirs:
        label = scale_dir.name
        if not (scale_dir / "DONE").exists():
            print(f"[skip] {label}: sem marcador DONE")
            continue

        manifest_path = scale_dir / "run_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
        n_drugs_scale = scale_drug_count(label, manifest)

        if n_drugs_scale is not None and n_drugs_scale <= args.min_drugs_scale:
            print(f"[skip] {label}: {n_drugs_scale} medicamentos "
                  f"(<= {args.min_drugs_scale}, fora do escopo pedido - ignorado)")
            continue

        rs_dir = scale_dir / "round_scores_raw"
        events_path = scale_dir / "events.json"
        scenario_path = scale_dir / "scenario_template.csv"
        if not rs_dir.exists() or not events_path.exists() or not scenario_path.exists():
            print(f"[skip] {label}: arquivos essenciais ausentes")
            continue

        round_scores = load_round_scores(rs_dir)
        events = load_events(str(events_path))
        scenario = load_scenario(str(scenario_path))
        enriched = join_round_scores_to_drugs(round_scores, events, scenario, drug_names)
        filtered, excluded, fallback_used = filter_low_volume_drugs(enriched, args.min_events_per_drug)

        evolution_rows = build_score_evolution_table(filtered)
        fraud_rows = build_fraud_cutoff_table(filtered, args.fraud_cutoff_pct)
        fraud_by_round = build_fraud_cutoff_table_by_round(filtered, args.fraud_cutoff_pct)

        summary_rows.append(compute_fraud_scale_summary(fraud_rows, label, n_drugs_scale))
        per_scale_data[label] = {
            "evolution": evolution_rows,
            "fraud": fraud_rows,
            "fraud_by_round": fraud_by_round,
        }
        print(f"[ok] {label}: {n_drugs_scale} medicamentos (escala), "
              f"{len(fraud_rows)} avaliados no relatorio de fraude" +
              (" (filtro de volume relaxado)" if fallback_used else ""))

    if not per_scale_data:
        print(f"Nenhuma escala com > {args.min_drugs_scale} medicamentos e marcador DONE "
              f"foi encontrada em {runroot}")
        sys.exit(1)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(summary_rows, per_scale_data, args.fraud_cutoff_pct, args.min_drugs_scale, output_path)
    print(f"\nPronto! {output_path}")


if __name__ == "__main__":
    main()

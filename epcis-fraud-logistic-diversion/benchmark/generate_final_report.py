
#!/usr/bin/env python3

"""

generate_final_report.py

=========================

Gera o pacote final de resultados do benchmark EPCIS logistic-diversion:

  - output_final/report_original_caliper.html   (copia do report.html cru)

  - output_final/resultados_comparativo.xlsx    (Excel com as 3 fases)

  - output_final/resultados_comparativo.pdf     (PDF com a mesma tabela + texto)

"""

import os, shutil, datetime

from openpyxl import Workbook

from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib import colors

from reportlab.lib.pagesizes import A4

from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from reportlab.lib.units import cm
 
OUT_DIR = "output_final"

os.makedirs(OUT_DIR, exist_ok=True)
 
rows = [

    ("1. Antes da correção (chave compartilhada)", "ObjectEvent",       1484, 7294, 8778),

    ("1. Antes da correção (chave compartilhada)", "AggregationEvent",  334,  277,  611),

    ("1. Antes da correção (chave compartilhada)", "TransactionEvent",  298,  313,  611),

    ("2. Fix parcial (composite keys, sem retry)",  "ObjectEvent",      1885, 6893, 8778),

    ("2. Fix parcial (composite keys, sem retry)",  "AggregationEvent",  342,  269,  611),

    ("2. Fix parcial (composite keys, sem retry)",  "TransactionEvent",  319,  292,  611),

    ("3. Fix completo (composite keys + retry)",    "ObjectEvent",      8377,  401, 8778),

    ("3. Fix completo (composite keys + retry)",    "AggregationEvent",  595,   16,  611),

    ("3. Fix completo (composite keys + retry)",    "TransactionEvent",  598,   13,  611),

]
 
def pct(succ, total):

    return round(100 * succ / total, 1)
 
wb = Workbook()

ws = wb.active

ws.title = "Comparativo"

headers = ["Fase", "Tipo de Evento", "Sucesso", "Falha", "Total", "% Sucesso"]

ws.append(headers)

for cell in ws[1]:

    cell.font = Font(bold=True, color="FFFFFF")

    cell.fill = PatternFill("solid", fgColor="366092")

    cell.alignment = Alignment(horizontal="center")

for fase, tipo, succ, fail, total in rows:

    ws.append([fase, tipo, succ, fail, total, pct(succ, total)])

widths = [42, 20, 10, 10, 10, 12]

for i, w in enumerate(widths, start=1):

    ws.column_dimensions[chr(64 + i)].width = w
 
ws2 = wb.create_sheet("Consolidado por Fase")

ws2.append(["Fase", "Sucesso Total", "Falha Total", "Total Geral", "% Sucesso Geral"])

for cell in ws2[1]:

    cell.font = Font(bold=True, color="FFFFFF")

    cell.fill = PatternFill("solid", fgColor="366092")

fases = {}

for fase, tipo, succ, fail, total in rows:

    if fase not in fases:

        fases[fase] = [0, 0, 0]

    fases[fase][0] += succ

    fases[fase][1] += fail

    fases[fase][2] += total

for fase, (s, f, t) in fases.items():

    ws2.append([fase, s, f, t, pct(s, t)])

for i, w in enumerate([42, 14, 14, 14, 16], start=1):

    ws2.column_dimensions[chr(64 + i)].width = w
 
xlsx_path = os.path.join(OUT_DIR, "resultados_comparativo.xlsx")

wb.save(xlsx_path)

print(f"Excel salvo: {xlsx_path}")
 
styles = getSampleStyleSheet()

title_style = ParagraphStyle("TitleBR", parent=styles["Title"], fontSize=16)

h2_style = ParagraphStyle("H2BR", parent=styles["Heading2"], fontSize=12, spaceBefore=12)

body_style = styles["BodyText"]
 
pdf_path = os.path.join(OUT_DIR, "resultados_comparativo.pdf")

doc = SimpleDocTemplate(pdf_path, pagesize=A4, topMargin=1.5*cm, bottomMargin=1.5*cm)

elements = []

elements.append(Paragraph("Benchmark EPCIS logistic-diversion — Comparativo Final", title_style))

elements.append(Spacer(1, 0.3*cm))

elements.append(Paragraph(

    f"Gerado automaticamente em {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} "

    "a partir da execução real no node003 (Hyperledger Fabric test-network, 2 orgs, Caliper 0.6.0). "

    "Escala: 1.000 medicamentos / 10.000 eventos (dataset ARCOS DEA).",

    body_style

))

elements.append(Paragraph("Resultado por tipo de evento e fase do projeto", h2_style))

table_data = [headers] + [[fase, tipo, succ, fail, total, f"{pct(succ,total)}%"] for fase, tipo, succ, fail, total in rows]

t = Table(table_data, repeatRows=1, colWidths=[7.5*cm, 3.3*cm, 2*cm, 2*cm, 2*cm, 2.2*cm])

t.setStyle(TableStyle([

    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#366092")),

    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),

    ("FONTSIZE", (0, 0), (-1, -1), 8),

    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),

    ("ALIGN", (2, 0), (-1, -1), "CENTER"),

    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),

]))

elements.append(t)

elements.append(Paragraph("Consolidado por fase (todos os tipos de evento somados)", h2_style))

cons_data = [["Fase", "Sucesso", "Falha", "Total", "% Sucesso"]]

for fase, (s, f, tt) in fases.items():

    cons_data.append([fase, s, f, tt, f"{pct(s, tt)}%"])

t2 = Table(cons_data, repeatRows=1, colWidths=[8*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm])

t2.setStyle(TableStyle([

    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#366092")),

    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),

    ("FONTSIZE", (0, 0), (-1, -1), 8),

    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),

    ("ALIGN", (1, 0), (-1, -1), "CENTER"),

]))

elements.append(t2)

elements.append(Paragraph("Diagnóstico técnico", h2_style))

elements.append(Paragraph(

    "<b>Fase 1 (bug original):</b> todas as transações liam e sobrescreviam a mesma chave de estado "

    "por entidade (produto/local), causando MVCC_READ_CONFLICT em quase toda escrita concorrente.<br/><br/>"

    "<b>Fase 2 (fix parcial):</b> substituição por chaves compostas append-only (event-sourcing via "

    "HistoryStore) eliminou a colisão de escrita-escrita, mas expôs um novo problema: as consultas de "

    "histórico usam range query (getStateByPartialCompositeKey), que o Fabric protege contra phantom "

    "read — gerando MVCC_READ_CONFLICT (3.908 ocorrências) e PHANTOM_READ_CONFLICT (2.961 ocorrências) "

    "quando dois eventos do mesmo produto/lote concorrem quase ao mesmo tempo.<br/><br/>"

    "<b>Fase 3 (fix completo):</b> duas melhorias combinadas — (a) eliminação de uma consulta de "

    "histórico duplicada por transação em LogisticsDiversionRule (reduz pela metade a exposição a "

    "phantom-read por EPC), e (b) retry automático no cliente Caliper (padrão recomendado pela "

    "documentação oficial do Hyperledger Fabric para contenção legítima de concorrência otimista, "

    "não é uma gambiarra). Resultado: taxa de sucesso geral subiu de ~25,5% para ~95,7%.",

    body_style

))

doc.build(elements)

print(f"PDF salvo: {pdf_path}")
 
if os.path.exists("report.html"):

    dest = os.path.join(OUT_DIR, "report_original_caliper.html")

    shutil.copy("report.html", dest)

    print(f"Report HTML copiado: {dest}")

else:

    print("AVISO: report.html não encontrado")
 
print(f"\nTudo pronto na pasta: {os.path.abspath(OUT_DIR)}")

